"""Summary Trip MHE — FastAPI backend (Vercel serverless, Postgres)."""
import io
import json
import logging
import os
import re
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import List, Optional

from fastapi import Depends, FastAPI, File, Header, HTTPException, Query, Request, UploadFile
from fastapi.responses import HTMLResponse, JSONResponse, StreamingResponse
from starlette.exceptions import HTTPException as StarletteHTTPException
from pydantic import BaseModel, Field

from . import auth, drive, gas, parsing, reconcile
from . import db as dbmod
from urllib.parse import urlsplit
from .db import conn, get_state, migrate
from .version import VERSION

log = logging.getLogger("mhe")
logging.basicConfig(level=logging.INFO)

BKK = timezone(timedelta(hours=7))
ROOT = Path(__file__).resolve().parent.parent
ADMIN_EMAIL = os.environ.get("ADMIN_EMAIL", "").strip().lower()
CRON_SECRET = os.environ.get("CRON_SECRET", "")
MAX_QTY = 100_000
MAX_UPLOAD = 4 * 1024 * 1024  # Vercel request body limit is 4.5 MB

app = FastAPI(title="Summary Trip MHE", docs_url=None, redoc_url=None)


# ======================================================================= errors
class AppError(HTTPException):
    def __init__(self, status: int, msg: str):
        super().__init__(status_code=status, detail=msg)


@app.exception_handler(StarletteHTTPException)   # also catches router 404/405
async def http_err(req: Request, exc: StarletteHTTPException):
    if exc.status_code == 404 and not isinstance(exc, AppError):
        return JSONResponse({"error": "ไม่พบ API นี้", "path": req.url.path}, status_code=404)
    if exc.status_code == 405:
        return JSONResponse({"error": "Method not allowed", "path": req.url.path}, status_code=405)
    return JSONResponse({"error": exc.detail}, status_code=exc.status_code)


@app.exception_handler(Exception)
async def any_err(_: Request, exc: Exception):
    log.exception("unhandled")
    return JSONResponse({"error": "ระบบขัดข้อง กรุณาลองใหม่"}, status_code=500)


# ====================================================================== startup
_booted = False


def boot():
    """Migrate + bootstrap admin once per cold start."""
    global _booted
    if _booted:
        return
    auth._secret()          # fail loudly if SECRET_KEY is missing
    migrate()
    with conn() as c:
        has_admin = c.execute("SELECT 1 FROM users WHERE role = 'admin' AND active LIMIT 1").fetchone()
        if not has_admin:
            email = ADMIN_EMAIL or "admin@mhe.local"
            existing = c.execute("SELECT id FROM users WHERE email = %s", (email,)).fetchone()
            if existing:
                c.execute("UPDATE users SET role = 'admin', active = TRUE WHERE id = %s", (existing["id"],))
                log.warning("BOOTSTRAP: promoted existing user %s to admin", email)
            else:
                pw = auth.random_password(14)
                c.execute(
                    """INSERT INTO users(email, display_name, password_hash, role, must_change_password)
                       VALUES (%s, 'Admin', %s, 'admin', TRUE)""", (email, auth.hash_password(pw)))
                # printed ONCE, only on the very first boot against an empty database
                log.warning("BOOTSTRAP ADMIN CREATED  email=%s  password=%s  (change it after first login)",
                            email, pw)
    _booted = True


def db():
    boot()
    with conn() as c:
        yield c


# ========================================================================= auth
def _cc_name(c, code):
    if not code:
        return ""
    r = c.execute("SELECT name FROM cost_centers WHERE code = %s", (code,)).fetchone()
    return r["name"] if r else ""


def _user_out(u, c):
    return {"id": u["id"], "email": u["email"], "display_name": u["display_name"],
            "employee_id": u.get("employee_id", ""),
            "cost_center": u["cost_center"], "cost_center_name": _cc_name(c, u["cost_center"]),
            "role": u["role"], "active": u["active"], "must_change_password": u["must_change_password"]}


def _valid_cc(c, code, *, required):
    """Cost Center must come from the admin-maintained list (when the list exists)."""
    code = (code or "").strip()
    has_list = c.execute("SELECT 1 FROM cost_centers WHERE active LIMIT 1").fetchone()
    if not code:
        if required and has_list:
            raise AppError(400, "กรุณาเลือก Cost Center")
        return ""
    if not c.execute("SELECT 1 FROM cost_centers WHERE code = %s AND active", (code,)).fetchone():
        raise AppError(400, f"ไม่พบ Cost Center {code} ในรายการ")
    return code


def current_user(authorization: str = Header(default=""), c=Depends(db, scope="function")):
    tok = authorization.removeprefix("Bearer ").strip()
    parsed = auth.read_token(tok) if tok else None
    if not parsed:
        raise AppError(401, "กรุณาเข้าสู่ระบบ")
    uid, ver = parsed
    u = c.execute("SELECT * FROM users WHERE id = %s", (uid,)).fetchone()
    if not u or u["token_version"] != ver:
        raise AppError(401, "เซสชันหมดอายุ กรุณาเข้าสู่ระบบใหม่")
    if not u["active"]:
        raise AppError(403, "บัญชีนี้ถูกปิดการใช้งาน")
    return u


def admin_user(u=Depends(current_user)):
    if u["role"] != "admin":
        raise AppError(403, "เฉพาะ Admin")
    return u


EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


class RegisterIn(BaseModel):
    email: str
    display_name: str
    employee_id: str = ""
    cost_center: str = ""
    password: str


EMAIL_DOMAIN = os.environ.get("DEFAULT_EMAIL_DOMAIN", "central.co.th").strip().lstrip("@").lower()
EMP_RE = re.compile(r"^[A-Za-z0-9-]{3,20}$")


def _full_email(s: str) -> str:
    """'trthirachet' -> 'trthirachet@central.co.th'; full addresses pass through."""
    s = (s or "").strip().lower()
    if s and "@" not in s and EMAIL_DOMAIN:
        s = f"{s}@{EMAIL_DOMAIN}"
    return s


def _valid_emp(c, emp: str, *, exclude_id=None) -> str:
    emp = (emp or "").strip().upper()
    if not EMP_RE.match(emp):
        raise AppError(400, "รหัสพนักงานใช้ตัวอักษร/ตัวเลข 3–20 ตัว")
    r = c.execute("SELECT id FROM users WHERE lower(employee_id) = lower(%s)", (emp,)).fetchone()
    if r and r["id"] != exclude_id:
        raise AppError(409, f"รหัสพนักงาน {emp} ถูกใช้สมัครแล้ว")
    return emp


class LoginIn(BaseModel):
    email: str
    password: str
    remember: bool = True


class PasswordIn(BaseModel):
    current: str = ""
    new: str


def _check_password(pw: str):
    if len(pw) < 8:
        raise AppError(400, "รหัสผ่านต้องมีอย่างน้อย 8 ตัวอักษร")
    if len(pw) > 128:
        raise AppError(400, "รหัสผ่านยาวเกินไป")


@app.post("/api/auth/register")
def register(body: RegisterIn, c=Depends(db, scope="function")):
    email = _full_email(body.email)
    name = body.display_name.strip()
    cc = body.cost_center.strip()
    if not EMAIL_RE.match(email):
        raise AppError(400, "รูปแบบอีเมลไม่ถูกต้อง")
    if not name or len(name) > 60:
        raise AppError(400, "กรุณากรอกชื่อที่แสดง (ไม่เกิน 60 ตัวอักษร)")
    cc = _valid_cc(c, cc, required=True)
    emp = _valid_emp(c, body.employee_id)
    _check_password(body.password)
    if c.execute("SELECT 1 FROM users WHERE email = %s", (email,)).fetchone():
        raise AppError(409, "อีเมลนี้ถูกใช้สมัครแล้ว")
    u = c.execute(
        """INSERT INTO users(email, display_name, employee_id, cost_center, password_hash)
           VALUES (%s,%s,%s,%s,%s) RETURNING *""",
        (email, name, emp, cc, auth.hash_password(body.password))).fetchone()
    return {"token": auth.make_token(u["id"], u["token_version"]), "user": _user_out(u, c)}


@app.post("/api/auth/login")
def login(body: LoginIn, c=Depends(db, scope="function")):
    ident = body.email.strip()
    u = None
    if ident and "@" not in ident:     # employee ID first, then short e-mail
        u = c.execute("SELECT * FROM users WHERE employee_id <> '' AND lower(employee_id) = lower(%s)",
                      (ident,)).fetchone()
    if not u:
        u = c.execute("SELECT * FROM users WHERE email = %s", (_full_email(ident),)).fetchone()
    if not u or not auth.verify_password(body.password, u["password_hash"]):
        raise AppError(401, "อีเมล/รหัสพนักงาน หรือรหัสผ่านไม่ถูกต้อง")
    if not u["active"]:
        raise AppError(403, "บัญชีนี้ถูกปิดการใช้งาน กรุณาติดต่อ Admin")
    c.execute("UPDATE users SET last_login_at = now() WHERE id = %s", (u["id"],))
    return {"token": auth.make_token(u["id"], u["token_version"], body.remember), "user": _user_out(u, c)}


@app.get("/api/me")
def me(u=Depends(current_user), c=Depends(db, scope="function")):
    return {"user": _user_out(u, c)}


@app.post("/api/me/password")
def change_password(body: PasswordIn, u=Depends(current_user), c=Depends(db, scope="function")):
    if not u["must_change_password"] and not auth.verify_password(body.current, u["password_hash"]):
        raise AppError(400, "รหัสผ่านปัจจุบันไม่ถูกต้อง")
    _check_password(body.new)
    r = c.execute(
        """UPDATE users SET password_hash = %s, must_change_password = FALSE,
               token_version = token_version + 1 WHERE id = %s RETURNING *""",
        (auth.hash_password(body.new), u["id"])).fetchone()
    return {"token": auth.make_token(r["id"], r["token_version"]), "user": _user_out(r, c)}


@app.post("/api/me/logout-all")
def logout_all(u=Depends(current_user), c=Depends(db, scope="function")):
    c.execute("UPDATE users SET token_version = token_version + 1 WHERE id = %s", (u["id"],))
    return {"ok": True}


# ========================================================================= meta
@app.get("/api/meta")
def meta(c=Depends(db, scope="function")):
    ccs = c.execute("SELECT code, name FROM cost_centers WHERE active ORDER BY code").fetchall()
    return {"version": VERSION, "cost_centers": ccs, "email_domain": EMAIL_DOMAIN}


# ========================================================================= sync
@app.post("/api/sync/auto")
def sync_auto(u=Depends(current_user)):
    if gas.configured():            # Apps Script pushes on its own schedule
        return {"ran": False, "reason": "push_mode"}
    return drive.sync(conn, force=False)


@app.get("/api/cron/sync")
def sync_cron(authorization: str = Header(default="")):
    if not CRON_SECRET or authorization != f"Bearer {CRON_SECRET}":
        raise AppError(401, "unauthorized")
    boot()
    if gas.configured():
        return gas.trigger() if gas.GAS_WEBAPP_URL else {"ran": False, "reason": "push_mode"}
    return drive.sync(conn, force=True)


# ================================================================== trip lookup
def _bkk_now():
    return datetime.now(BKK)


@app.get("/api/docs/search")
def doc_search(q: str = "", u=Depends(current_user), c=Depends(db, scope="function")):
    raw = q.strip().upper().replace(" ", "")
    if not raw:
        return {"items": []}
    key = parsing.doc_key(raw)
    zero_only = raw.strip("0") == ""
    like_raw = raw.replace("%", "").replace("_", r"\_") + "%"
    like_key = key.replace("%", "").replace("_", r"\_") + "%"
    since = _bkk_now().date() - timedelta(days=45)

    plan = c.execute(
        f"""SELECT load_no, load_key, max(trip_no) trip_no, max(truck_id) truck_id,
                   count(*) stores, max(plan_date) plan_date
            FROM plan_lines
            WHERE (upper(load_no) LIKE %s OR upper(trip_no) LIKE %s
                   {'' if zero_only else 'OR load_key LIKE %s OR trip_key LIKE %s'})
              AND (plan_date IS NULL OR plan_date >= %s)
            GROUP BY load_no, load_key ORDER BY max(plan_date) DESC NULLS LAST, load_no DESC LIMIT 12""",
        [like_raw, like_raw] + ([] if zero_only else [like_key, like_key]) + [since]).fetchall()
    recs = c.execute(
        f"""SELECT doc_no, doc_key, leg, complete, trip_no, truck_id
            FROM trip_records
            WHERE upper(doc_no) LIKE %s {'' if zero_only else 'OR doc_key LIKE %s'}
            ORDER BY updated_at DESC LIMIT 30""",
        [like_raw] + ([] if zero_only else [like_key])).fetchall()

    legs = {}
    for r in recs:
        legs.setdefault(r["doc_key"], {})[r["leg"]] = r
    items, seen = [], set()
    for p in plan:
        k = p["load_key"]
        seen.add(k)
        items.append({"doc_no": p["load_no"], "trip_no": p["trip_no"], "truck_id": p["truck_id"],
                      "stores": p["stores"], "source": "plan", "legs": _legs(legs.get(k, {}))})
    for r in recs:
        if r["doc_key"] in seen:
            continue
        seen.add(r["doc_key"])
        items.append({"doc_no": r["doc_no"], "trip_no": r["trip_no"], "truck_id": r["truck_id"],
                      "stores": None, "source": "record", "legs": _legs(legs.get(r["doc_key"], {}))})
    return {"items": items[:15]}


def _legs(d):
    return {leg: ("complete" if r["complete"] else "incomplete") for leg, r in d.items()}


def _plan_for(c, key):
    rows = c.execute(
        "SELECT * FROM plan_lines WHERE load_key = %s ORDER BY seq, id", (key,)).fetchall()
    if not rows:
        loads = c.execute("SELECT DISTINCT load_key FROM plan_lines WHERE trip_key = %s", (key,)).fetchall()
        if len(loads) == 1:
            rows = c.execute("SELECT * FROM plan_lines WHERE load_key = %s ORDER BY seq, id",
                             (loads[0]["load_key"],)).fetchall()
    return rows


def _record(c, key, leg):
    r = c.execute(
        """SELECT r.*, cu.display_name AS created_by_name, uu.display_name AS updated_by_name
           FROM trip_records r
           LEFT JOIN users cu ON cu.id = r.created_by LEFT JOIN users uu ON uu.id = r.updated_by
           WHERE r.leg = %s AND r.doc_key = %s""", (leg, key)).fetchone()
    if not r:
        return None, []
    lines = c.execute("SELECT * FROM record_lines WHERE record_id = %s ORDER BY sort_order, id",
                      (r["id"],)).fetchall()
    return r, lines


def _qty(l):
    return {"pallet": l["pallet"], "totebox": l["totebox"], "rollcage": l["rollcage"], "box": l["box"]}


@app.get("/api/docs/{doc}")
def doc_detail(doc: str, leg: str = Query(..., pattern="^(out|ret)$"),
               u=Depends(current_user), c=Depends(db, scope="function")):
    doc = doc.strip()
    key = parsing.doc_key(doc)
    if not key:
        raise AppError(400, "กรุณาระบุ Document No.")
    plan = _plan_for(c, key)
    rec, rec_lines = _record(c, key, leg)
    base, base_lines = (_record(c, key, "out") if leg == "ret" else (None, []))

    rows, index = [], {}

    def add(code, name, in_plan):
        if code in index:
            r = rows[index[code]]
            r["in_plan"] = r["in_plan"] or in_plan
            if not r["store_name"] and name:
                r["store_name"] = name
            return r
        r = {"store_code": code, "store_name": name, "in_plan": in_plan, "baseline": None,
             "saved": None, "deleted": False, "comment": ""}
        index[code] = len(rows)
        rows.append(r)
        return r

    for p in plan:
        add(p["store_code"], p["store_name"], True)
    for l in base_lines:
        r = add(l["store_code"], l["store_name"], l["in_plan"])
        r["baseline"] = {"pallet": 0, "totebox": 0, "rollcage": 0, "box": 0} if l["deleted"] else _qty(l)
    for l in rec_lines:
        r = add(l["store_code"], l["store_name"], l["in_plan"])
        r["saved"] = _qty(l)
        r["deleted"] = l["deleted"]
        r["comment"] = l["comment"]

    head = plan[0] if plan else None
    src = rec or base
    return {
        "doc_no": (rec and rec["doc_no"]) or (head and head["load_no"]) or (base and base["doc_no"]) or doc,
        "load_no": (head and head["load_no"]) or (src and src["load_no"]) or "",
        "trip_no": (head and head["trip_no"]) or (src and src["trip_no"]) or "",
        "truck_id": (head and head["truck_id"]) or (src and src["truck_id"]) or "",
        "truck_type": (head and head["truck_type"]) or (src and src["truck_type"]) or "",
        "transporter": (head and head["transporter"]) or (src and src["transporter"]) or "",
        "in_plan": bool(plan),
        "has_baseline": base is not None,
        "leg": leg,
        "record": None if not rec else {
            "id": rec["id"], "door_no": rec["door_no"], "complete": rec["complete"],
            "updated_at": rec["updated_at"].isoformat(), "updated_by": rec["updated_by_name"],
            "created_by": rec["created_by_name"],
        },
        "baseline_door": base["door_no"] if base else None,
        "rows": rows,
    }


@app.get("/api/stores")
def store_search(q: str = "", u=Depends(current_user), c=Depends(db, scope="function")):
    q = q.strip()
    if not q:
        return {"items": []}
    like = "%" + q.replace("%", "").replace("_", r"\_") + "%"
    rows = c.execute(
        """SELECT code, name, bu FROM stores WHERE code LIKE %s OR name ILIKE %s
           ORDER BY (code = %s) DESC, length(code), code LIMIT 20""",
        (q + "%" if q.isdigit() else like, like, parsing.store_code(q))).fetchall()
    return {"items": rows}


# ====================================================================== records
class LineIn(BaseModel):
    store_code: str
    store_name: str = ""
    in_plan: bool = False
    deleted: bool = False
    entered: bool = False
    pallet: Optional[int] = None
    totebox: Optional[int] = None
    rollcage: Optional[int] = None
    box: Optional[int] = None
    comment: str = ""


class RecordIn(BaseModel):
    doc_no: str
    leg: str = Field(pattern="^(out|ret)$")
    door_no: Optional[int] = None
    base_updated_at: Optional[str] = None   # optimistic concurrency
    lines: List[LineIn]


QTY = ("pallet", "totebox", "rollcage", "box")
QTY_LABEL = {"pallet": "Pallet", "totebox": "Totebox", "rollcage": "Rollcage", "box": "Box"}


def _clean_line(l: LineIn):
    code = parsing.store_code(l.store_code)
    if not code or len(code) > 20:
        raise AppError(400, f"รหัสสาขาไม่ถูกต้อง: {l.store_code}")
    vals = {}
    for k in QTY:
        v = getattr(l, k)
        if v is not None and (v < 0 or v > MAX_QTY):
            raise AppError(400, f"สาขา {code}: {QTY_LABEL[k]} ต้องอยู่ระหว่าง 0–{MAX_QTY}")
        vals[k] = 0 if (l.deleted or v is None) else int(v)
    entered = l.deleted or l.entered or any(getattr(l, k) is not None for k in QTY) or bool(l.comment.strip())
    return code, {
        "store_name": l.store_name.strip()[:120], "in_plan": l.in_plan, "deleted": l.deleted,
        "comment": l.comment.strip()[:500], "entered": entered, **vals,
    }


@app.post("/api/records")
def save_record(body: RecordIn, u=Depends(current_user), c=Depends(db, scope="function")):
    doc = body.doc_no.strip()
    key = parsing.doc_key(doc)
    if not key or len(doc) > 40:
        raise AppError(400, "Document No. ไม่ถูกต้อง")
    if body.door_no is None or not (1 <= body.door_no <= 999):
        raise AppError(400, "กรุณาระบุ Door No. (1–999)")

    lines, order = {}, []
    for l in body.lines:
        code, d = _clean_line(l)
        if code in lines:
            raise AppError(400, f"สาขา {code} ซ้ำในรายการ")
        lines[code] = d
        order.append(code)
    stored = {k: v for k, v in lines.items() if v["entered"]}
    if not stored:
        raise AppError(400, "ยังไม่ได้กรอกข้อมูลสาขาใดเลย")

    plan = _plan_for(c, key)
    head = plan[0] if plan else None
    expected = {p["store_code"] for p in plan}
    if body.leg == "ret":
        b, bl = _record(c, key, "out")
        expected |= {l["store_code"] for l in bl}
    else:
        b = None
    expected |= set(lines)
    complete = expected <= set(stored)

    # lock the row (if any) for the concurrency check
    old = c.execute("SELECT * FROM trip_records WHERE doc_key = %s AND leg = %s FOR UPDATE",
                    (key, body.leg)).fetchone()
    if old and body.base_updated_at != old["updated_at"].isoformat():
        who = c.execute("SELECT display_name FROM users WHERE id = %s", (old["updated_by"],)).fetchone()
        raise AppError(409, f"Trip นี้เพิ่งถูกแก้ไขโดย {who['display_name'] if who else 'ผู้อื่น'} "
                            "กรุณาโหลดข้อมูลล่าสุดก่อนบันทึก")
    if not old and body.base_updated_at:
        raise AppError(409, "ข้อมูลนี้ถูกลบหรือเปลี่ยนแปลง กรุณาโหลดใหม่")

    load_no = (head and head["load_no"]) or (b and b["load_no"]) or (old and old["load_no"]) or doc
    trip_no = (head and head["trip_no"]) or (b and b["trip_no"]) or (old and old["trip_no"]) or ""
    truck = (head and head["truck_id"]) or (b and b["truck_id"]) or (old and old["truck_id"]) or ""
    ttype = (head and head["truck_type"]) or (b and b["truck_type"]) or (old and old["truck_type"]) or ""
    tport = (head and head["transporter"]) or (b and b["transporter"]) or (old and old["transporter"]) or ""
    fields = (body.door_no, bool(plan), complete, len(expected), len(stored), u["id"],
              load_no, parsing.doc_key(load_no), trip_no, parsing.doc_key(trip_no), truck, ttype, tport)

    if old:
        old_lines = {l["store_code"]: l for l in c.execute(
            "SELECT * FROM record_lines WHERE record_id = %s", (old["id"],)).fetchall()}
        changes = _diff(old, old_lines, body.door_no, stored)
        rec = c.execute(
            """UPDATE trip_records SET door_no=%s, in_plan=%s, complete=%s, expected_stores=%s,
                   entered_stores=%s, updated_by=%s, updated_at=now(), load_no=%s, load_key=%s,
                   trip_no=%s, trip_key=%s, truck_id=%s, truck_type=%s, transporter=%s
               WHERE id=%s RETURNING *""", fields + (old["id"],)).fetchone()
        c.execute("DELETE FROM record_lines WHERE record_id = %s", (rec["id"],))
        action = "update"
    else:
        rec = c.execute(
            """INSERT INTO trip_records(door_no, in_plan, complete, expected_stores, entered_stores,
                   updated_by, load_no, load_key, trip_no, trip_key, truck_id, truck_type, transporter,
                   created_by, doc_no, doc_key, leg)
               VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
               ON CONFLICT (doc_key, leg) DO NOTHING RETURNING *""",
            fields + (u["id"], doc, key, body.leg)).fetchone()
        if not rec:   # someone else created it a moment ago
            raise AppError(409, "Trip นี้เพิ่งถูกบันทึกโดยผู้อื่น กรุณาโหลดข้อมูลล่าสุดก่อนบันทึก")
        changes = [f"บันทึกครั้งแรก · Door {body.door_no} · {len(stored)} สาขา"]
        action = "create"

    with c.cursor() as cur:
        cur.executemany(
            """INSERT INTO record_lines(record_id, store_code, store_name, in_plan, deleted,
                   pallet, totebox, rollcage, box, comment, sort_order)
               VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)""",
            [(rec["id"], code, d["store_name"], d["in_plan"], d["deleted"], d["pallet"], d["totebox"],
              d["rollcage"], d["box"], d["comment"], order.index(code)) for code, d in stored.items()])
        names = [(code, d["store_name"]) for code, d in stored.items() if d["store_name"]]
        cur.executemany(
            """INSERT INTO stores(code, name) VALUES (%s,%s) ON CONFLICT (code) DO NOTHING""", names)
    if changes:
        c.execute("INSERT INTO record_history(record_id, user_id, action, changes) VALUES (%s,%s,%s,%s)",
                  (rec["id"], u["id"], action, json.dumps(changes, ensure_ascii=False)))
    return {"id": rec["id"], "complete": complete, "expected": len(expected), "entered": len(stored),
            "updated_at": rec["updated_at"].isoformat(), "changed": bool(changes)}


def _diff(old, old_lines, door, new_lines):
    ch = []
    if old["door_no"] != door:
        ch.append(f"Door {old['door_no']} → {door}")
    for code, n in new_lines.items():
        o = old_lines.get(code)
        if not o:
            ch.append(f"เพิ่มสาขา {code}" + (" (ลบ = 0)" if n["deleted"] else ""))
            continue
        if n["deleted"] and not o["deleted"]:
            ch.append(f"ลบสาขา {code} (บันทึกเป็น 0)")
            continue
        if o["deleted"] and not n["deleted"]:
            ch.append(f"ยกเลิกการลบสาขา {code}")
        for k in QTY:
            if o[k] != n[k]:
                ch.append(f"{code} {QTY_LABEL[k]} {o[k]} → {n[k]}")
        if (o["comment"] or "") != n["comment"]:
            ch.append(f"{code} Comment: {n['comment'] or '(ลบ)'}")
    for code in old_lines:
        if code not in new_lines:
            ch.append(f"ล้างข้อมูลสาขา {code}")
    return ch


def _list_filters(date_from, date_to, q, status, leg):
    where, args = ["(r.created_at AT TIME ZONE 'Asia/Bangkok')::date BETWEEN %s AND %s"], [date_from, date_to]
    if leg in ("out", "ret"):
        where.append("r.leg = %s")
        args.append(leg)
    if status == "complete":
        where.append("r.complete")
    elif status == "incomplete":
        where.append("NOT r.complete")
    if q.strip():
        like = "%" + q.strip().replace("%", "").replace("_", r"\_") + "%"
        k = parsing.doc_key(q)
        where.append("""(r.doc_no ILIKE %s OR r.trip_no ILIKE %s OR r.truck_id ILIKE %s OR r.doc_key LIKE %s
                         OR EXISTS (SELECT 1 FROM record_lines x WHERE x.record_id = r.id
                                    AND (x.store_code LIKE %s OR x.store_name ILIKE %s)))""")
        args += [like, like, like, "%" + k + "%", q.strip() + "%", like]
    return " AND ".join(where), args


def _parse_date(s, default):
    if not s:
        return default
    try:
        return date.fromisoformat(s)
    except ValueError:
        raise AppError(400, "รูปแบบวันที่ไม่ถูกต้อง")


LIST_SQL = """
SELECT r.id, r.doc_no, r.leg, r.trip_no, r.truck_id, r.truck_type, r.transporter, r.door_no, r.complete, r.in_plan,
       r.expected_stores, r.entered_stores, r.created_at, r.updated_at,
       cu.display_name AS created_by, uu.display_name AS updated_by,
       COALESCE(SUM(l.pallet),0)::int pallet, COALESCE(SUM(l.totebox),0)::int totebox,
       COALESCE(SUM(l.rollcage),0)::int rollcage, COALESCE(SUM(l.box),0)::int box
FROM trip_records r
LEFT JOIN record_lines l ON l.record_id = r.id
LEFT JOIN users cu ON cu.id = r.created_by LEFT JOIN users uu ON uu.id = r.updated_by
WHERE {where}
GROUP BY r.id, cu.display_name, uu.display_name
ORDER BY r.updated_at DESC
"""


@app.get("/api/records")
def list_records(date_from: str = "", date_to: str = "", q: str = "", status: str = "all",
                 leg: str = "all", u=Depends(current_user), c=Depends(db, scope="function")):
    today = _bkk_now().date()
    df, dt = _parse_date(date_from, today), _parse_date(date_to, today)
    where, args = _list_filters(df, dt, q, "all", leg)
    rows = c.execute(LIST_SQL.format(where=where) + " LIMIT 500", args).fetchall()
    counts = {"all": len(rows), "complete": sum(r["complete"] for r in rows)}
    counts["incomplete"] = counts["all"] - counts["complete"]
    if status == "complete":
        rows = [r for r in rows if r["complete"]]
    elif status == "incomplete":
        rows = [r for r in rows if not r["complete"]]
    return {"items": [_rec_out(r) for r in rows], "counts": counts}


def _rec_out(r):
    d = dict(r)
    for k in ("created_at", "updated_at"):
        d[k] = r[k].isoformat()
    return d


@app.get("/api/records/export")
def export_records(date_from: str = "", date_to: str = "", q: str = "", status: str = "all",
                   leg: str = "all", u=Depends(current_user), c=Depends(db, scope="function")):
    import openpyxl
    today = _bkk_now().date()
    df, dt = _parse_date(date_from, today), _parse_date(date_to, today)
    where, args = _list_filters(df, dt, q, status, leg)
    rows = c.execute(
        f"""SELECT r.updated_at, r.created_at, r.doc_no, r.leg, r.trip_no, r.truck_id, r.door_no,
                   r.complete, cu.display_name created_by, uu.display_name updated_by,
                   uu.employee_id, uu.cost_center, cc.name AS cost_center_name, r.truck_type, r.transporter, l.store_code, l.store_name, l.in_plan, l.deleted,
                   l.pallet, l.totebox, l.rollcage, l.box, l.comment
            FROM trip_records r JOIN record_lines l ON l.record_id = r.id
            LEFT JOIN users cu ON cu.id = r.created_by LEFT JOIN users uu ON uu.id = r.updated_by
            LEFT JOIN cost_centers cc ON cc.code = uu.cost_center
            WHERE {where} ORDER BY r.updated_at DESC, l.sort_order""", args).fetchall()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Records"
    ws.append(["Created", "Updated", "Document No.", "ขา", "Trip No.", "Truck", "Door", "สถานะ",
               "ผู้บันทึก", "แก้ไขล่าสุดโดย", "รหัสพนักงาน", "Cost Center", "Cost Center Name", "ประเภทรถ", "Transporter", "Store Code", "Store Name", "ในแผน",
               "ลบ (=0)", "Pallet", "Totebox", "Rollcage", "Box", "Comment"])
    for r in rows:
        ws.append([_xl_dt(r["created_at"]), _xl_dt(r["updated_at"]), r["doc_no"],
                   "ขากลับ" if r["leg"] == "ret" else "ขาออก", r["trip_no"], r["truck_id"], r["door_no"],
                   "ครบ" if r["complete"] else "ยังไม่ครบ", r["created_by"], r["updated_by"], r["employee_id"], r["cost_center"], r["cost_center_name"], r["truck_type"], r["transporter"],
                   r["store_code"], r["store_name"], "Y" if r["in_plan"] else "N", "Y" if r["deleted"] else "",
                   r["pallet"], r["totebox"], r["rollcage"], r["box"], r["comment"]])
    return _xlsx(wb, f"SummaryTrip_{df:%Y%m%d}-{dt:%Y%m%d}.xlsx")


def _xl_dt(v):
    return v.astimezone(BKK).replace(tzinfo=None) if v else None


def _xlsx(wb, name):
    for ws in wb.worksheets:
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 8), 40)
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": _content_disposition(name)})


def _content_disposition(name):
    from urllib.parse import quote
    ascii_name = name.encode("ascii", "replace").decode().replace("?", "_")
    return f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quote(name)}"


@app.get("/api/records/{rid}")
def record_detail(rid: int, u=Depends(current_user), c=Depends(db, scope="function")):
    r = c.execute(LIST_SQL.format(where="r.id = %s"), (rid,)).fetchone()
    if not r:
        raise AppError(404, "ไม่พบข้อมูล")
    lines = c.execute("SELECT * FROM record_lines WHERE record_id = %s ORDER BY sort_order, id",
                      (rid,)).fetchall()
    hist = c.execute(
        """SELECT h.at, h.action, h.changes, u.display_name FROM record_history h
           LEFT JOIN users u ON u.id = h.user_id WHERE h.record_id = %s ORDER BY h.at DESC, h.id DESC""",
        (rid,)).fetchall()
    return {
        "record": _rec_out(r),
        "lines": [{k: l[k] for k in ("store_code", "store_name", "in_plan", "deleted", "pallet",
                                     "totebox", "rollcage", "box", "comment")} for l in lines],
        "history": [{"at": h["at"].isoformat(), "action": h["action"], "changes": h["changes"],
                     "by": h["display_name"]} for h in hist],
    }


# ======================================================================== admin
@app.get("/api/admin/import")
def import_status(u=Depends(admin_user), c=Depends(db, scope="function")):
    files = c.execute(
        """SELECT name, modified_time, imported_at, status, row_count, trip_count, error
           FROM plan_files ORDER BY imported_at DESC LIMIT 40""").fetchall()
    today = _bkk_now().date()
    stats = c.execute(
        """SELECT count(*) FILTER (WHERE (imported_at AT TIME ZONE 'Asia/Bangkok')::date = %s) files_today,
                  count(*) FILTER (WHERE status = 'error' AND (imported_at AT TIME ZONE 'Asia/Bangkok')::date = %s) errors_today
           FROM plan_files""", (today, today)).fetchone()
    loads = c.execute("SELECT count(DISTINCT load_key) n FROM plan_lines WHERE plan_date = %s",
                      (today,)).fetchone()["n"]
    sa_email = ""
    try:
        sa_email = json.loads(os.environ.get("GOOGLE_SERVICE_ACCOUNT_JSON", "{}")).get("client_email", "")
    except json.JSONDecodeError:
        pass
    return {
        "configured": drive.configured() or gas.configured(),
        "mode": "apps_script" if gas.configured() else ("service_account" if drive.configured() else ""),
        **gas.status(c),
        "service_account": sa_email,
        "auto_minutes": drive.AUTO_SYNC_MINUTES,
        "last_sync_at": get_state(c, "last_sync_at"), "last_run_at": get_state(c, "last_run_at"),
        "last_error": get_state(c, "last_sync_error", ""),
        "files_today": stats["files_today"], "errors_today": stats["errors_today"], "loads_today": loads,
        "files": [{**f, "modified_time": f["modified_time"] and f["modified_time"].isoformat(),
                   "imported_at": f["imported_at"].isoformat()} for f in files],
    }


@app.post("/api/admin/sync")
def sync_now(u=Depends(admin_user)):
    if gas.configured():
        return gas.trigger()
    if drive.configured():
        return drive.sync(conn, force=True)
    return {"ran": False, "reason": "not_configured"}


# ------------------------------------------------------- Apps Script push mode
class PushFile(BaseModel):
    id: str
    name: str = ""
    modified: Optional[str] = None


class PushCheckIn(BaseModel):
    files: List[PushFile]


class PushPlanIn(PushFile):
    content_b64: str


class PushDoneIn(BaseModel):
    seen: int = 0
    pushed: List[dict] = []
    errors: List[str] = []
    pending: int = 0


def _push_auth(x_push_secret: str = Header(default="")):
    if not gas.configured():
        raise AppError(503, "PLAN_PUSH_SECRET ยังไม่ได้ตั้งค่าบน Vercel (ต้องยาวอย่างน้อย 16 ตัวอักษร)")
    if not gas.secret_ok(x_push_secret):
        raise AppError(401, "รหัส PUSH_SECRET ไม่ตรงกับที่ตั้งไว้บน Vercel")
    return True


@app.post("/api/push/plan/check")
def push_check(body: PushCheckIn, ok=Depends(_push_auth), c=Depends(db, scope="function")):
    files = [f.model_dump() for f in body.files[:500]]
    return {"needed": gas.needed(c, files)}


@app.post("/api/push/plan")
def push_plan(body: PushPlanIn, ok=Depends(_push_auth), c=Depends(db, scope="function")):
    import base64
    import binascii
    try:
        content = base64.b64decode(body.content_b64, validate=True)
    except (binascii.Error, ValueError):
        raise AppError(400, "content_b64 ไม่ถูกต้อง")
    if len(content) > MAX_UPLOAD:
        raise AppError(413, "ไฟล์ใหญ่เกิน 4 MB")
    mt = gas.ts(body.modified) or datetime.now(timezone.utc)
    res = drive.import_plan_content(c, drive_file_id=body.id, name=body.name or body.id,
                                    modified_time=mt, content=content)
    return res


@app.post("/api/push/plan/done")
def push_done(body: PushDoneIn, ok=Depends(_push_auth), c=Depends(db, scope="function")):
    gas.record_run(c, body.model_dump())
    return {"ok": True}


async def _read_upload(f: UploadFile):
    data = await f.read()
    if len(data) > MAX_UPLOAD:
        raise AppError(413, "ไฟล์ใหญ่เกิน 4 MB")
    if not (f.filename or "").lower().endswith((".xlsx", ".xlsm")):
        raise AppError(400, "รองรับเฉพาะไฟล์ .xlsx")
    return data


@app.post("/api/admin/plan/upload")
async def plan_upload(file: UploadFile = File(...), u=Depends(admin_user)):
    import hashlib
    data = await _read_upload(file)
    boot()
    with conn() as c:
        res = drive.import_plan_content(
            c, drive_file_id="upload:" + hashlib.sha1(data).hexdigest(), name=file.filename,
            modified_time=datetime.now(timezone.utc), content=data)
    if res["status"] != "ok":
        raise AppError(400, res["error"])
    return res


class UserPatch(BaseModel):
    role: Optional[str] = Field(default=None, pattern="^(user|admin)$")
    active: Optional[bool] = None
    cost_center: Optional[str] = None
    display_name: Optional[str] = None
    employee_id: Optional[str] = None


@app.get("/api/admin/users")
def users(u=Depends(admin_user), c=Depends(db, scope="function")):
    rows = c.execute(
        """SELECT u.id, u.email, u.display_name, u.employee_id, u.cost_center, COALESCE(cc.name, '') AS cost_center_name,
                  u.role, u.active, u.created_at, u.last_login_at
           FROM users u LEFT JOIN cost_centers cc ON cc.code = u.cost_center
           ORDER BY u.active DESC, u.role, u.display_name""").fetchall()
    return {"items": [{**r, "created_at": r["created_at"].isoformat(),
                       "last_login_at": r["last_login_at"] and r["last_login_at"].isoformat()} for r in rows]}


@app.patch("/api/admin/users/{uid}")
def patch_user(uid: int, body: UserPatch, u=Depends(admin_user), c=Depends(db, scope="function")):
    t = c.execute("SELECT * FROM users WHERE id = %s FOR UPDATE", (uid,)).fetchone()
    if not t:
        raise AppError(404, "ไม่พบผู้ใช้")
    demote = body.role == "user" and t["role"] == "admin"
    disable = body.active is False and t["active"]
    if (demote or disable) and t["role"] == "admin":
        n = c.execute("SELECT count(*) n FROM users WHERE role = 'admin' AND active").fetchone()["n"]
        if n <= 1:
            raise AppError(400, "ต้องมี Admin ที่ใช้งานได้อย่างน้อย 1 คน")
    sets, args = [], []
    if body.role:
        sets.append("role = %s")
        args.append(body.role)
    if body.active is not None:
        sets.append("active = %s")
        args.append(body.active)
        if not body.active:
            sets.append("token_version = token_version + 1")   # kick existing sessions
    if body.cost_center is not None:
        sets.append("cost_center = %s")
        args.append(_valid_cc(c, body.cost_center, required=False))
    if body.employee_id is not None:
        sets.append("employee_id = %s")
        args.append(_valid_emp(c, body.employee_id, exclude_id=uid))
    if body.display_name is not None and body.display_name.strip():
        sets.append("display_name = %s")
        args.append(body.display_name.strip()[:60])
    if not sets:
        return {"ok": True}
    c.execute(f"UPDATE users SET {', '.join(sets)} WHERE id = %s", args + [uid])
    return {"ok": True}


@app.post("/api/admin/users/{uid}/reset-password")
def reset_password(uid: int, u=Depends(admin_user), c=Depends(db, scope="function")):
    pw = auth.random_password(10)
    r = c.execute(
        """UPDATE users SET password_hash = %s, must_change_password = TRUE,
               token_version = token_version + 1 WHERE id = %s RETURNING email""",
        (auth.hash_password(pw), uid)).fetchone()
    if not r:
        raise AppError(404, "ไม่พบผู้ใช้")
    return {"email": r["email"], "temp_password": pw}


# ------------------------------------------------------------------ cost centers
class CostCenterIn(BaseModel):
    code: str
    name: str = ""


class CostCenterPatch(BaseModel):
    name: Optional[str] = None
    active: Optional[bool] = None


CC_RE = re.compile(r"^[A-Za-z0-9._-]{1,20}$")


@app.get("/api/admin/cost-centers")
def cc_list(u=Depends(admin_user), c=Depends(db, scope="function")):
    rows = c.execute(
        """SELECT cc.code, cc.name, cc.active, count(u.id)::int AS users
           FROM cost_centers cc LEFT JOIN users u ON u.cost_center = cc.code
           GROUP BY cc.code ORDER BY cc.active DESC, cc.code""").fetchall()
    orphans = c.execute(
        """SELECT u.cost_center AS code, count(*)::int AS users FROM users u
           WHERE u.cost_center <> '' AND NOT EXISTS (SELECT 1 FROM cost_centers cc WHERE cc.code = u.cost_center)
           GROUP BY 1 ORDER BY 1""").fetchall()
    return {"items": rows, "unlisted": orphans}


@app.post("/api/admin/cost-centers")
def cc_create(body: CostCenterIn, u=Depends(admin_user), c=Depends(db, scope="function")):
    code, name = body.code.strip(), body.name.strip()[:80]
    if not CC_RE.match(code):
        raise AppError(400, "รหัส Cost Center ใช้ได้เฉพาะตัวอักษร/ตัวเลข ไม่เกิน 20 ตัว")
    if not name:
        raise AppError(400, "กรุณาตั้งชื่อ Cost Center")
    r = c.execute("INSERT INTO cost_centers(code, name) VALUES (%s,%s) ON CONFLICT DO NOTHING RETURNING code",
                  (code, name)).fetchone()
    if not r:
        raise AppError(409, f"มี Cost Center {code} อยู่แล้ว")
    return {"ok": True}


@app.patch("/api/admin/cost-centers/{code}")
def cc_patch(code: str, body: CostCenterPatch, u=Depends(admin_user), c=Depends(db, scope="function")):
    if not c.execute("SELECT 1 FROM cost_centers WHERE code = %s", (code,)).fetchone():
        raise AppError(404, "ไม่พบ Cost Center")
    if body.name is not None:
        if not body.name.strip():
            raise AppError(400, "กรุณาตั้งชื่อ Cost Center")
        c.execute("UPDATE cost_centers SET name = %s WHERE code = %s", (body.name.strip()[:80], code))
    if body.active is not None:
        c.execute("UPDATE cost_centers SET active = %s WHERE code = %s", (body.active, code))
    return {"ok": True}


@app.delete("/api/admin/cost-centers/{code}")
def cc_delete(code: str, u=Depends(admin_user), c=Depends(db, scope="function")):
    n = c.execute("SELECT count(*)::int n FROM users WHERE cost_center = %s", (code,)).fetchone()["n"]
    if n:
        raise AppError(400, f"มีผู้ใช้ {n} คนอยู่ใน Cost Center นี้ — ใช้ปิดการใช้งานแทน หรือย้ายผู้ใช้ก่อน")
    c.execute("DELETE FROM cost_centers WHERE code = %s", (code,))
    return {"ok": True}


# -------------------------------------------------------------------- reconcile
@app.post("/api/admin/mhe/upload")
async def mhe_upload(file: UploadFile = File(...), u=Depends(admin_user)):
    data = await _read_upload(file)
    try:
        lines, st = parsing.parse_mhe(data)
    except parsing.ParseError as e:
        raise AppError(400, str(e))
    boot()
    with conn() as c:
        up = c.execute(
            """INSERT INTO mhe_uploads(filename, uploaded_by, row_count, trip_count, date_min, date_max)
               VALUES (%s,%s,%s,%s,%s,%s) RETURNING id""",
            (file.filename, u["id"], st["rows"], st["trips"], st["date_min"], st["date_max"])).fetchone()
        with c.cursor() as cur:
            cur.executemany(
                """INSERT INTO mhe_lines(upload_id, trip_no, trip_key, store_code, store_name, bu, mtype,
                       tn_date, key_date, qty_in, qty_out)
                   VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
                   ON CONFLICT (trip_key, store_code, mtype) DO UPDATE SET
                       upload_id = EXCLUDED.upload_id, trip_no = EXCLUDED.trip_no,
                       store_name = EXCLUDED.store_name, bu = EXCLUDED.bu, tn_date = EXCLUDED.tn_date,
                       key_date = EXCLUDED.key_date, qty_in = EXCLUDED.qty_in, qty_out = EXCLUDED.qty_out""",
                [(up["id"], l["trip_no"], l["trip_key"], l["store_code"], l["store_name"], l["bu"],
                  l["mtype"], l["tn_date"], l["key_date"], l["qty_in"], l["qty_out"]) for l in lines])
            st_rows = {}
            for l in lines:
                st_rows[l["store_code"]] = (l["store_name"], l["bu"])
            cur.executemany(
                """INSERT INTO stores(code, name, bu) VALUES (%s,%s,%s)
                   ON CONFLICT (code) DO UPDATE SET bu = EXCLUDED.bu,
                       name = CASE WHEN stores.name = '' THEN EXCLUDED.name ELSE stores.name END""",
                [(k, v[0], v[1]) for k, v in st_rows.items()])
    return {"rows": st["rows"], "lines": st["lines"], "trips": st["trips"], "skipped": st["skipped"],
            "date_min": st["date_min"] and st["date_min"].isoformat(),
            "date_max": st["date_max"] and st["date_max"].isoformat()}


@app.get("/api/admin/mhe/uploads")
def mhe_uploads(u=Depends(admin_user), c=Depends(db, scope="function")):
    rows = c.execute(
        """SELECT m.id, m.filename, m.uploaded_at, m.row_count, m.trip_count, m.date_min, m.date_max,
                  u.display_name AS by FROM mhe_uploads m LEFT JOIN users u ON u.id = m.uploaded_by
           ORDER BY m.uploaded_at DESC LIMIT 10""").fetchall()
    rng = c.execute("SELECT min(tn_date) a, max(tn_date) b FROM mhe_lines").fetchone()
    return {"items": [{**r, "uploaded_at": r["uploaded_at"].isoformat(),
                       "date_min": r["date_min"] and r["date_min"].isoformat(),
                       "date_max": r["date_max"] and r["date_max"].isoformat()} for r in rows],
            "range": {"min": rng["a"] and rng["a"].isoformat(), "max": rng["b"] and rng["b"].isoformat()}}


def _recon_args(date_from, date_to, leg, mtype):
    if leg not in reconcile.LEG_COL:
        raise AppError(400, "leg ไม่ถูกต้อง")
    if mtype not in ("all", "pallet", "totebox"):
        raise AppError(400, "type ไม่ถูกต้อง")
    today = _bkk_now().date()
    df = _parse_date(date_from, today - timedelta(days=6))
    dt = _parse_date(date_to, today)
    if df > dt:
        raise AppError(400, "ช่วงวันที่ไม่ถูกต้อง")
    return df, dt


def _ser(v):
    return v.isoformat() if isinstance(v, (date, datetime)) else v


MODES = ("both", "all")


@app.get("/api/admin/reconcile")
def reconcile_view(date_from: str = "", date_to: str = "", leg: str = "ret", bu: str = "",
                   mtype: str = "all", mode: str = "both",
                   u=Depends(admin_user), c=Depends(db, scope="function")):
    df, dt = _recon_args(date_from, date_to, leg, mtype)
    if mode not in MODES:
        raise AppError(400, "mode ไม่ถูกต้อง")
    res = reconcile.run(c, date_from=df, date_to=dt, leg=leg, bu=bu, mtype=mtype, mode=mode)
    res.pop("all_rows")
    res["diffs_total"] = len(res["diffs"])
    res["diffs"] = [{k: _ser(v) for k, v in d.items()} for d in res["diffs"][:5000]]
    res["only_mhe_trips"] = [{k: _ser(v) for k, v in d.items()} for d in res["only_mhe_trips"][:200]]
    res["range"] = {"from": df.isoformat(), "to": dt.isoformat()}
    res["categories"] = c.execute(
        "SELECT id, name, active FROM recon_categories ORDER BY active DESC, sort_order, id").fetchall()
    return res


@app.get("/api/admin/reconcile/export")
def reconcile_export(date_from: str = "", date_to: str = "", leg: str = "ret", bu: str = "",
                     mtype: str = "all", mode: str = "both",
                     u=Depends(admin_user), c=Depends(db, scope="function")):
    import openpyxl
    df, dt = _recon_args(date_from, date_to, leg, mtype)
    if mode not in MODES:
        raise AppError(400, "mode ไม่ถูกต้อง")
    res = reconcile.run(c, date_from=df, date_to=dt, leg=leg, bu=bu, mtype=mtype, mode=mode)
    col = reconcile.MHE_COL_LABEL[leg]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Compare"
    ws.append(["TNDate / วันที่คีย์", "Trip No.", "Document No.", "BU", "Store Code", "Store Name", "Type",
               "LP", f"MHE ({col})", "ต่าง (LP-MHE)", "สถานะ", "คีย์โดย",
               "ประเภทปัญหา", "หมายเหตุ", "แก้ไขแล้ว", "บันทึกหมายเหตุโดย"])
    for r in sorted(res["all_rows"], key=lambda r: (str(r["tn_date"]), r["trip_no"], r["store_code"])):
        ws.append([r["tn_date"], r["trip_no"], r["doc_no"], r["bu"], r["store_code"], r["store_name"],
                   reconcile.TYPE_LABEL[r["mtype"]], r["lp"], r["mhe"], r["diff"],
                   reconcile.STATUS_LABEL[r["status"]], r["keyed_by"],
                   r["category"], r["note"], "Y" if r["resolved"] else "", r["note_by"]])
    ws4 = wb.create_sheet("By category")
    ws4.append(["ประเภทปัญหา", "จำนวนรายการ"])
    for x in res["by_category"]:
        ws4.append([x["name"], x["count"]])
    ws2 = wb.create_sheet("MHE only (LP ไม่ได้คีย์)")
    ws2.append(["TNDate", "Trip No.", "BU", "จำนวนสาขา", f"รวม {col}"])
    for t in res["only_mhe_trips"]:
        ws2.append([t["tn_date"], t["trip_no"], t["bu"], t["stores"], t["qty"]])
    ws3 = wb.create_sheet("By person")
    ws3.append(["ผู้คีย์", "ตรงกัน", "ทั้งหมด", "%"])
    for p in res["people"]:
        ws3.append([p["name"], p["match"], p["total"], p["pct"]])
    tag = "all" if mode == "all" else "both-keyed"
    return _xlsx(wb, f"Reconcile_{leg}_{tag}_{df:%Y%m%d}-{dt:%Y%m%d}.xlsx")


class ReconNoteIn(BaseModel):
    leg: str = Field(pattern="^(out|ret)$")
    trip_key: str
    store_code: str
    mtype: str = Field(pattern="^(pallet|totebox)$")
    category_id: Optional[int] = None
    note: str = ""
    resolved: bool = False


@app.put("/api/admin/reconcile/note")
def reconcile_note(body: ReconNoteIn, u=Depends(admin_user), c=Depends(db, scope="function")):
    if not body.trip_key.strip() or not body.store_code.strip():
        raise AppError(400, "ข้อมูลรายการไม่ครบ")
    if body.category_id is not None and not c.execute(
            "SELECT 1 FROM recon_categories WHERE id = %s", (body.category_id,)).fetchone():
        raise AppError(400, "ไม่พบประเภทปัญหา")
    note = body.note.strip()[:1000]
    if body.category_id is None and not note and not body.resolved:
        c.execute("DELETE FROM recon_notes WHERE leg=%s AND trip_key=%s AND store_code=%s AND mtype=%s",
                  (body.leg, body.trip_key, body.store_code, body.mtype))
        return {"ok": True, "cleared": True}
    c.execute(
        """INSERT INTO recon_notes(leg, trip_key, store_code, mtype, category_id, note, resolved, updated_by)
           VALUES (%s,%s,%s,%s,%s,%s,%s,%s)
           ON CONFLICT (leg, trip_key, store_code, mtype) DO UPDATE SET
               category_id = EXCLUDED.category_id, note = EXCLUDED.note, resolved = EXCLUDED.resolved,
               updated_by = EXCLUDED.updated_by, updated_at = now()""",
        (body.leg, body.trip_key, body.store_code, body.mtype, body.category_id, note, body.resolved, u["id"]))
    return {"ok": True}


class CategoryIn(BaseModel):
    name: str


class CategoryPatch(BaseModel):
    name: Optional[str] = None
    active: Optional[bool] = None


@app.get("/api/admin/recon-categories")
def recon_categories(u=Depends(admin_user), c=Depends(db, scope="function")):
    return {"items": c.execute(
        """SELECT rc.id, rc.name, rc.active, count(n.*)::int AS used
           FROM recon_categories rc LEFT JOIN recon_notes n ON n.category_id = rc.id
           GROUP BY rc.id ORDER BY rc.active DESC, rc.sort_order, rc.id""").fetchall()}


@app.post("/api/admin/recon-categories")
def recon_category_add(body: CategoryIn, u=Depends(admin_user), c=Depends(db, scope="function")):
    name = body.name.strip()[:60]
    if not name:
        raise AppError(400, "กรุณาตั้งชื่อประเภทปัญหา")
    r = c.execute(
        """INSERT INTO recon_categories(name, sort_order)
           VALUES (%s, (SELECT COALESCE(max(sort_order), 0) + 10 FROM recon_categories))
           ON CONFLICT (name) DO NOTHING RETURNING id""", (name,)).fetchone()
    if not r:
        raise AppError(409, f"มีประเภท \"{name}\" อยู่แล้ว")
    return {"id": r["id"]}


@app.patch("/api/admin/recon-categories/{cid}")
def recon_category_patch(cid: int, body: CategoryPatch, u=Depends(admin_user), c=Depends(db, scope="function")):
    if not c.execute("SELECT 1 FROM recon_categories WHERE id = %s", (cid,)).fetchone():
        raise AppError(404, "ไม่พบประเภทปัญหา")
    if body.name is not None:
        name = body.name.strip()[:60]
        if not name:
            raise AppError(400, "กรุณาตั้งชื่อประเภทปัญหา")
        if c.execute("SELECT 1 FROM recon_categories WHERE name = %s AND id <> %s", (name, cid)).fetchone():
            raise AppError(409, f"มีประเภท \"{name}\" อยู่แล้ว")
        c.execute("UPDATE recon_categories SET name = %s WHERE id = %s", (name, cid))
    if body.active is not None:
        c.execute("UPDATE recon_categories SET active = %s WHERE id = %s", (body.active, cid))
    return {"ok": True}


@app.delete("/api/admin/recon-categories/{cid}")
def recon_category_delete(cid: int, u=Depends(admin_user), c=Depends(db, scope="function")):
    n = c.execute("SELECT count(*)::int n FROM recon_notes WHERE category_id = %s", (cid,)).fetchone()["n"]
    if n:
        raise AppError(400, f"มีรายการใช้ประเภทนี้อยู่ {n} รายการ — ใช้ปิดการใช้งานแทน")
    c.execute("DELETE FROM recon_categories WHERE id = %s", (cid,))
    return {"ok": True}


# ======================================================================= health
@app.get("/api/health")
def health(request: Request):
    """Deployment check. Shows WHICH settings exist, never their values."""
    sk = os.environ.get("SECRET_KEY", "")
    out = {
        "version": VERSION, "path": request.url.path,
        "env": {
            "DATABASE_URL": bool(dbmod.DATABASE_URL),
            "DATABASE_URL_from": dbmod.DB_ENV_NAME or "(none)",
            "DATABASE_host": _db_host(),
            "DATABASE_URL_pooled": "-pooler" in dbmod.DATABASE_URL,
            "SECRET_KEY_ok": len(sk) >= 32,
            "ADMIN_EMAIL": bool(ADMIN_EMAIL),
            "CRON_SECRET": bool(CRON_SECRET),
            "DRIVE_configured": drive.configured(),
            "PUSH_configured": gas.configured(),
            "GAS_WEBAPP_URL": bool(gas.GAS_WEBAPP_URL),
            "VERCEL_REGION": os.environ.get("VERCEL_REGION", ""),
        },
    }
    try:
        boot()
        with conn() as c:
            out["db"] = "ok"
            out["users"] = c.execute("SELECT count(*)::int n FROM users").fetchone()["n"]
    except Exception as e:  # noqa: BLE001
        msg = str(e)
        pw = urlsplit(dbmod.DATABASE_URL).password if dbmod.DATABASE_URL else None
        if pw:
            msg = msg.replace(pw, "***")
        out["db"] = f"error: {type(e).__name__}: {msg[:300]}"
    return out


def _db_host():
    try:
        return urlsplit(dbmod.DATABASE_URL).hostname or ""
    except ValueError:
        return "(invalid URL)"


# ======================================================================== pages
@app.get("/", response_class=HTMLResponse)
@app.get("/index.html", response_class=HTMLResponse)
def index():
    html = (ROOT / "static" / "index.html").read_text(encoding="utf-8")
    return HTMLResponse(html.replace("__VERSION__", VERSION),
                        headers={"Cache-Control": "no-cache"})


@app.get("/manifest.webmanifest")
def manifest():
    return JSONResponse({
        "name": "Summary Trip MHE", "short_name": "Trip MHE", "start_url": "/", "display": "standalone",
        "background_color": "#F4F2EE", "theme_color": "#16302F", "lang": "th",
        "icons": [{"src": "/icon.svg", "sizes": "any", "type": "image/svg+xml"}],
    }, media_type="application/manifest+json")


@app.get("/icon.svg")
def icon():
    svg = ('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 64 64"><rect width="64" height="64" rx="14" '
           'fill="#D9822B"/><g fill="none" stroke="#16302F" stroke-width="4" stroke-linecap="round" '
           'stroke-linejoin="round"><path d="M10 18h26v20H10z M36 26h9l7 7v5H36z"/><circle cx="19" cy="44" r="5"/>'
           '<circle cx="44" cy="44" r="5"/></g></svg>')
    return HTMLResponse(svg, media_type="image/svg+xml", headers={"Cache-Control": "public, max-age=86400"})


# SPA fallback: any non-API GET path serves the app (e.g. /login, /index, deep links)
@app.get("/{full_path:path}", include_in_schema=False)
def spa_fallback(full_path: str):
    last = full_path.rsplit("/", 1)[-1]
    if full_path.startswith("api/") or full_path == "api" or "." in last:   # API or a file like favicon.ico
        raise StarletteHTTPException(404)
    return index()
