"""Excel parsing for the plan-load report and the MHE Daily Movement file.

openpyxl only (no pandas) to keep the Vercel bundle and cold start small.
"""
import io
import re
from datetime import date, datetime

import openpyxl


class ParseError(Exception):
    pass


def doc_key(s) -> str:
    """Normalised key for Trip/Load/Document numbers.

    Excel often turns '001449276' into the number 1449276, so digits-only
    values are compared without leading zeros.
    """
    s = cell_str(s).upper().replace(" ", "")
    if s.isdigit():
        s = s.lstrip("0") or "0"
    return s


LOAD_NO_DIGITS = 9   # MHE shows Load No. as 001461288; Excel stores 1461288


def load_no_fmt(s: str) -> str:
    s = s.strip()
    if s.isdigit() and len(s) < LOAD_NO_DIGITS:
        return s.zfill(LOAD_NO_DIGITS)
    return s


def store_code(v) -> str:
    s = cell_str(v)
    if s.isdigit():
        return str(int(s))
    return s.upper()


def cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        return str(v)
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def to_int(v) -> int:
    if v is None or v == "":
        return 0
    try:
        return int(round(float(str(v).replace(",", "").strip())))
    except ValueError:
        return 0


def to_num(v):
    if v is None or v == "":
        return None
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        return None


def to_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def _norm_header(v) -> str:
    return re.sub(r"[^a-z0-9]", "", cell_str(v).lower())


def _load(content: bytes):
    try:
        return openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise ParseError(f"เปิดไฟล์ Excel ไม่ได้: {e}") from e


def _find_header(rows, must_have, scan=15):
    for i, row in enumerate(rows[:scan]):
        norm = [_norm_header(c) for c in row]
        if all(any(m == h for h in norm) for m in must_have):
            return i, norm
    raise ParseError(
        "ไม่พบแถวหัวตารางที่มีคอลัมน์ " + ", ".join(f'"{m}"' for m in must_have)
    )


# ------------------------------------------------------------------ plan load
def parse_plan(content: bytes):
    """Summary plan load daily report -> list of dict lines.

    Layout (from the original Streamlit app): header on row 3 with columns
    NO. | Trip No. | ID Truck | Store Code | Store  Name | Trip No. (=Load No.) |
    Pallet | Rollcage | Boxes. Trip/Truck/Load are merged cells -> forward-fill.
    """
    wb = _load(content)
    last_err = None
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        try:
            h, norm = _find_header(rows, ["storecode", "tripno"])
        except ParseError as e:
            last_err = e
            continue
        trip_cols = [i for i, n in enumerate(norm) if n == "tripno"]
        col = {
            "trip": trip_cols[0],
            "load": trip_cols[1] if len(trip_cols) > 1 else None,
            "truck": next((i for i, n in enumerate(norm) if n in ("idtruck", "truckid", "truck")), None),
            "code": norm.index("storecode"),
            "name": next((i for i, n in enumerate(norm) if n == "storename"), None),
            "pallet": next((i for i, n in enumerate(norm) if n == "pallet"), None),
            "rollcage": next((i for i, n in enumerate(norm) if n in ("rollcage", "roll")), None),
            "boxes": next((i for i, n in enumerate(norm) if n in ("boxes", "box")), None),
        }

        def g(row, key):
            i = col[key]
            return row[i] if i is not None and i < len(row) else None

        out = []
        cur_trip = cur_truck = cur_load = ""
        seq_by_load = {}
        for row in rows[h + 1:]:
            if not any(c not in (None, "") for c in row):
                continue
            t = cell_str(g(row, "trip"))
            if t and t != cur_trip:
                # new trip group: don't carry the previous group's load/truck over
                cur_trip, cur_truck, cur_load = t, "", ""
            tr = cell_str(g(row, "truck"))
            if tr and not cur_truck:
                cur_truck = tr
            # The Load No. column sometimes repeats the Trip No. on the 2nd row of a
            # group (real file 25.09.26) -> only the first real value of a group counts.
            ld = cell_str(g(row, "load"))
            if ld and not cur_load and ld != cur_trip:
                cur_load = load_no_fmt(ld)
            code = store_code(g(row, "code"))
            if not code or not re.search(r"\d", code):
                continue  # totals / notes
            load_no = cur_load or cur_trip
            if not load_no:
                continue
            seq_by_load[load_no] = seq_by_load.get(load_no, 0) + 1
            out.append({
                "load_no": load_no,
                "trip_no": cur_trip,
                "truck_id": cur_truck,
                "store_code": code,
                "store_name": cell_str(g(row, "name")),
                "seq": seq_by_load[load_no],
                "plan_pallet": to_num(g(row, "pallet")),
                "plan_rollcage": to_num(g(row, "rollcage")),
                "plan_boxes": to_num(g(row, "boxes")),
            })
        if out:
            return out
        last_err = ParseError("พบหัวตารางแต่ไม่มีรายการสาขา")
    raise last_err or ParseError("ไฟล์ว่าง")


# ------------------------------------------------------------------ MHE file
MHE_TYPES = {"pallet": "pallet", "totebox": "totebox", "tote": "totebox"}


def parse_mhe(content: bytes):
    """AllBUDailyMovement export -> (aggregated lines, stats)."""
    wb = _load(content)
    ws = wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    h, norm = _find_header(rows, ["tripno", "storecode", "type", "in", "out"])
    ix = {n: i for i, n in enumerate(norm) if n}

    def g(row, name):
        i = ix.get(name)
        return row[i] if i is not None and i < len(row) else None

    agg = {}
    raw_rows = skipped = 0
    dmin = dmax = None
    for row in rows[h + 1:]:
        if not any(c not in (None, "") for c in row):
            continue
        trip = cell_str(g(row, "tripno"))
        code = store_code(g(row, "storecode"))
        mtype = MHE_TYPES.get(_norm_header(g(row, "type")))
        if not trip or not code:
            continue
        raw_rows += 1
        if not mtype:
            skipped += 1
            continue
        tn = to_date(g(row, "tndate"))
        kd = to_date(g(row, "keydate"))
        if tn:
            dmin = tn if dmin is None or tn < dmin else dmin
            dmax = tn if dmax is None or tn > dmax else dmax
        k = (doc_key(trip), code, mtype)
        a = agg.get(k)
        if not a:
            a = agg[k] = {
                "trip_no": trip, "trip_key": k[0], "store_code": code, "mtype": mtype,
                "store_name": cell_str(g(row, "storename")), "bu": cell_str(g(row, "bu")),
                "tn_date": tn, "key_date": kd, "qty_in": 0, "qty_out": 0,
            }
        a["qty_in"] += to_int(g(row, "in"))
        a["qty_out"] += to_int(g(row, "out"))
        if tn and (a["tn_date"] is None or tn > a["tn_date"]):
            a["tn_date"] = tn
    if not agg:
        raise ParseError("ไม่พบข้อมูล Pallet / Tote Box ในไฟล์")
    lines = list(agg.values())
    stats = {
        "rows": raw_rows, "skipped": skipped, "lines": len(lines),
        "trips": len({l["trip_key"] for l in lines}), "date_min": dmin, "date_max": dmax,
    }
    return lines, stats


# ------------------------------------------------------------- store master
def parse_store_master(content: bytes):
    """'Master Store' sheet of the plan file -> [(code, name, bu)]. Empty if absent."""
    try:
        wb = _load(content)
    except ParseError:
        return []
    ws = next((w for w in wb.worksheets if _norm_header(w.title) == "masterstore"), None)
    if ws is None:
        return []
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    try:
        h, norm = _find_header(rows, ["storecode", "storename"], scan=5)
    except ParseError:
        return []
    ic, iname = norm.index("storecode"), norm.index("storename")
    ibu = norm.index("bu") if "bu" in norm else None
    out = {}
    for r in rows[h + 1:]:
        code = store_code(r[ic] if ic < len(r) else None)
        if not code or not code.isdigit():
            continue
        name = cell_str(r[iname] if iname < len(r) else None)
        bu = cell_str(r[ibu]) if ibu is not None and ibu < len(r) else ""
        out[code] = (code, name, bu)
    return list(out.values())
